"""P30 approval-gated artifact generation from persisted Mission records only."""
from __future__ import annotations
import hashlib,json,re
from pathlib import Path
from sqlalchemy import func,select
from sqlalchemy.exc import OperationalError
from app.models.ai_mission import AIMission,AIMissionEvent
from app.models.artifact import Artifact,ArtifactEvidence,ArtifactVersion
from app.models.file_asset import FileAsset
from app.models.mission_file_source import MissionFileSource
from app.models.paper import Paper
from app.models.paper_chunk import PaperChunk
from app.models.connector import ArtifactDataSource, DataSource, MissionDataSource
from app.services.database import PROJECT_ROOT,SessionLocal,initialize_database

class ArtifactError(ValueError):pass
class ArtifactService:
    TYPES={"RESEARCH_BRIEF":"pdf","SOLUTION_DOCUMENT":"pdf","PRESENTATION":"pptx","DATA_REPORT":"xlsx","DELIVERY_PACKAGE":"pdf"}
    def __init__(self,session_factory=SessionLocal,*,initialize=True,root=None):
        if initialize:initialize_database()
        self.s=session_factory;self.root=root or PROJECT_ROOT/"work"/"artifacts"
    def generate(self,mission_id,artifact_type):
        if artifact_type not in self.TYPES:raise ArtifactError("不支持的 Artifact 类型。")
        s=self.s()
        try:
            mission=s.get(AIMission,mission_id)
            if not mission:raise ArtifactError("Mission 不存在。")
            if mission.status=="FAILED":raise ArtifactError("失败的 Mission 不能生成正式 Artifact。")
            row=s.scalar(select(Artifact).where(Artifact.mission_id==mission_id,Artifact.artifact_type==artifact_type))
            if row and row.generation_rounds>=3: row.status="WAITING_ADAPTIVE_REVIEW";s.commit();raise ArtifactError("已达到 3 次生成上限，等待人工 Adaptive Review。")
            version=(row.version+1) if row else 1
            evidence=self._evidence(s,mission)
            content=self._content(mission,artifact_type,evidence)
            if not row:
                row=Artifact(mission_id=mission_id,project_id=mission.solution_project_id,artifact_type=artifact_type,title=content["title"],version=version,status="GENERATING",source_type="MIXED" if any(x["evidence_type"]=="CUSTOMER_PROVIDED_DOCUMENT" for x in evidence) else "AI_GENERATED");s.add(row);s.flush()
            row.version=version;row.title=content["title"];row.status="GENERATING";row.generation_rounds+=1;row.content_summary=content["summary"];row.evidence_count=len(evidence);row.evidence_coverage=100 if evidence else 0
            filename=self._filename(row); path=self.root/filename;path.parent.mkdir(parents=True,exist_ok=True);self._write(path,row,content,evidence)
            row.file_path=str(path);row.status="GENERATED";s.query(ArtifactEvidence).filter_by(artifact_id=row.id).delete()
            for item in evidence:s.add(ArtifactEvidence(artifact_id=row.id,**item))
            # P31 retains provenance without copying enterprise rows into the Artifact.
            try:
                for link in s.scalars(select(MissionDataSource).where(MissionDataSource.mission_id==mission_id)).all():
                    if not s.scalar(select(ArtifactDataSource).where(ArtifactDataSource.artifact_id==row.id,ArtifactDataSource.data_source_id==link.data_source_id)):
                        s.add(ArtifactDataSource(artifact_id=row.id,data_source_id=link.data_source_id))
            except OperationalError:
                # Legacy isolated fixtures and older local databases can still
                # generate a draft without P31 provenance tables.
                pass
            digest=hashlib.sha256(path.read_bytes()).hexdigest();s.add(ArtifactVersion(artifact_id=row.id,version=version,change_summary="Initial generation" if version==1 else "Regenerated after human revision",content_hash=digest))
            validation=self._validate(s,row,evidence);row.status="NEEDS_REVIEW" if validation["status"]!="FAIL" else "FAILED";self._trace(s,mission_id,"GENERATE_ARTIFACT",row,"Artifact generated as reviewable draft.");self._trace(s,mission_id,"VALIDATE_ARTIFACT",row,validation["summary"]);s.commit();return self.detail(row.id,s)
        finally:s.close()
    def review(self,artifact_id,status,comment=""):
        if status not in {"APPROVED","REVISION_REQUESTED","REJECTED"}:raise ArtifactError("无效审核状态。")
        s=self.s()
        try:
            row=self._row(s,artifact_id)
            if row.status!="NEEDS_REVIEW":raise ArtifactError("Artifact 尚未进入人工审核。")
            row.status=status;self._trace(s,row.mission_id,"ARTIFACT_REVIEW",row,comment or f"Human review: {status}");s.commit();return self.detail(row.id,s)
        finally:s.close()
    def revise(self,artifact_id,reason):
        s=self.s()
        try:
            row=self._row(s,artifact_id)
            if row.status!="REVISION_REQUESTED":raise ArtifactError("只有请求修订的 Artifact 可以生成新版本。")
            mission_id=row.mission_id;kind=row.artifact_type;self._trace(s,mission_id,"REQUEST_REVISION",row,reason or "Reviewer requested revision.");s.commit()
        finally:s.close()
        return self.generate(mission_id,kind)
    def list(self,mission_id):
        s=self.s()
        try:return [self._data(x,s) for x in s.scalars(select(Artifact).where(Artifact.mission_id==mission_id).order_by(Artifact.updated_at.desc())).all()]
        finally:s.close()
    def detail(self,artifact_id,s=None):
        close=s is None;s=s or self.s()
        try:return self._data(self._row(s,artifact_id),s,True)
        finally:
            if close:s.close()
    def analytics(self):
        s=self.s()
        try:
            rows=list(s.scalars(select(Artifact)).all());total=len(rows)
            return {"total_artifacts":total,"generated":sum(x.status=="NEEDS_REVIEW" for x in rows),"approved":sum(x.status=="APPROVED" for x in rows),"revision_requested":sum(x.status=="REVISION_REQUESTED" for x in rows),"validation_failed":sum(x.status=="FAILED" for x in rows),"average_evidence_coverage":round(sum(x.evidence_coverage for x in rows)/total,1) if total else 0,"average_revision_count":round(sum(max(x.version-1,0) for x in rows)/total,1) if total else 0}
        finally:s.close()
    def download_path(self,artifact_id):
        s=self.s()
        try:
            row=self._row(s,artifact_id)
            if row.status!="APPROVED":raise ArtifactError("Artifact 尚未通过人工审核，不能作为正式交付下载。")
            path=Path(row.file_path).resolve()
            if self.root.resolve() not in path.parents or not path.is_file():raise ArtifactError("Artifact 文件不存在或不在受控工作区。")
            return path
        finally:s.close()
    def _evidence(self,s,mission):
        refs=self._decode(mission.evidence_refs_json);out=[]
        for item in refs:
            paper=s.get(Paper,item.get("paper_id"));chunk=s.get(PaperChunk,item.get("chunk_id")) if item.get("chunk_id") else None
            if paper and (not item.get("chunk_id") or chunk):out.append({"evidence_type":"KNOWLEDGE_BASE","paper_id":paper.paper_id,"chunk_id":item.get("chunk_id"),"source":paper.title,"section":item.get("section","") or "","claim_summary":"Retrieved Mission Evidence"})
        for link in s.scalars(select(MissionFileSource).where(MissionFileSource.mission_id==mission.id)).all():
            asset=s.get(FileAsset,link.file_id)
            if asset:out.append({"evidence_type":"CUSTOMER_PROVIDED_DOCUMENT","paper_id":None,"chunk_id":None,"source":asset.filename,"section":"Customer Material","claim_summary":"Customer-provided material; requires confirmation."})
        return out
    def _content(self,m,kind,evidence):
        sources=[x["source"] for x in evidence];base={"Title":m.title,"Research Question":m.goal or "NEEDS_CONFIRMATION","Evidence":", ".join(sources) if sources else "NEEDS_CONFIRMATION","Risks / Limitations":"AI GENERATED DRAFT · NEEDS_CONFIRMATION","Next Steps":"Human review required before delivery."}
        if kind=="DATA_REPORT":base={"Summary":"NO_REAL_DATA_SOURCE","Metrics":"No verified enterprise dataset attached.","Data Dictionary":"NEEDS_CONFIRMATION","Analysis":"No data analysis was fabricated.","Notes":"Attach a real approved data source before analysis."}
        if kind=="SOLUTION_DOCUMENT":base.update({"Customer Problem":"NEEDS_CONFIRMATION","Business Requirements":"NEEDS_CONFIRMATION","AI Solution":"AI GENERATED DRAFT · NEEDS_CONFIRMATION","System Architecture":"Use existing ResearchOS workflow after confirmation.","Implementation Roadmap":"Confirm scope → review → delivery.","Security & Risk":"Customer material remains isolated from knowledge-base Evidence.","Acceptance Criteria":"TO_BE_VALIDATED"})
        if kind=="PRESENTATION":base.update({"Customer Problem":"NEEDS_CONFIRMATION","Requirements":"NEEDS_CONFIRMATION","Solution Architecture":"Mission → Evidence → Human Review → Delivery","Expected Outcome":"TO_BE_VALIDATED"})
        summary="AI GENERATED DRAFT. Evidence and customer material are explicitly distinguished; human review is required."
        if not evidence: summary+=" NEEDS_CONFIRMATION: no verified Evidence is attached; no research conclusion is included."
        return {"title":f"{kind.replace('_',' ').title()} · {m.title}","sections":base,"summary":summary}
    def _filename(self,row):return f"{row.id}_v{row.version}.{self.TYPES[row.artifact_type]}"
    def _write(self,path,row,content,evidence):
        if path.suffix==".pptx":
            from pptx import Presentation
            deck=Presentation();
            for heading,value in content["sections"].items():slide=deck.slides.add_slide(deck.slide_layouts[1]);slide.shapes.title.text=heading;slide.placeholders[1].text=str(value)
            deck.save(path);return
        if path.suffix==".xlsx":
            from openpyxl import Workbook
            book=Workbook();sheet=book.active;sheet.title="Summary";sheet.append(["Field","Value"])
            for key,value in content["sections"].items():sheet.append([key,str(value)])
            book.save(path);return
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfgen import canvas
        pdf=canvas.Canvas(str(path),pagesize=A4);w,h=A4;y=h-48;pdf.setFont("Helvetica-Bold",15);pdf.drawString(42,y,content["title"][:92]);y-=30;pdf.setFont("Helvetica",9)
        for key,value in content["sections"].items():
            for line in [f"{key}: {value}"[i:i+105] for i in range(0,len(f"{key}: {value}"),105)]:
                if y<48:pdf.showPage();y=h-48;pdf.setFont("Helvetica",9)
                pdf.drawString(42,y,line);y-=14
        pdf.drawString(42,max(y,36),f"Sources: {', '.join(x['source'] for x in evidence) or 'NEEDS_CONFIRMATION'}");pdf.save()
        if row.artifact_type=="SOLUTION_DOCUMENT":
            from docx import Document
            doc=Document();doc.add_heading(content["title"],0)
            for key,value in content["sections"].items():doc.add_heading(key,1);doc.add_paragraph(str(value))
            doc.save(path.with_suffix(".docx"))
    def _validate(self,s,row,evidence):
        if not Path(row.file_path).is_file() or not row.content_summary:return {"status":"FAIL","summary":"Artifact file is missing or empty."}
        if not evidence:return {"status":"WARNING","summary":"No verified Evidence attached; draft remains NEEDS_CONFIRMATION."}
        return {"status":"PASS","summary":"Artifact file and persisted Evidence references validated."}
    @staticmethod
    def _decode(value):
        try:return json.loads(value or "[]")
        except json.JSONDecodeError:return []
    @staticmethod
    def _row(s,ident):
        row=s.get(Artifact,ident)
        if not row:raise ArtifactError("Artifact 不存在。")
        return row
    def _data(self,row,s,detail=False):
        data={"id":row.id,"mission_id":row.mission_id,"project_id":row.project_id,"artifact_type":row.artifact_type,"title":row.title,"version":row.version,"status":row.status,"source_type":row.source_type,"content_summary":row.content_summary,"evidence_count":row.evidence_count,"evidence_coverage":row.evidence_coverage,"generation_rounds":row.generation_rounds,"created_at":row.created_at,"updated_at":row.updated_at}
        if detail:
            data.update({"preview":{"title":row.title,"summary":row.content_summary,"review_status":row.status},"versions":[{"version":x.version,"change_summary":x.change_summary,"created_at":x.created_at} for x in s.scalars(select(ArtifactVersion).where(ArtifactVersion.artifact_id==row.id).order_by(ArtifactVersion.version)).all()],"evidence":[{"evidence_type":x.evidence_type,"paper_id":x.paper_id,"chunk_id":x.chunk_id,"source":x.source,"section":x.section,"claim_summary":x.claim_summary} for x in s.scalars(select(ArtifactEvidence).where(ArtifactEvidence.artifact_id==row.id)).all()]})
            try:
                data["data_sources"]=[{"id":source.id,"name":source.name,"type":source.source_type} for link in s.scalars(select(ArtifactDataSource).where(ArtifactDataSource.artifact_id==row.id)).all() if (source:=s.get(DataSource,link.data_source_id))]
            except OperationalError:
                data["data_sources"]=[]
        return data
    @staticmethod
    def _trace(s,mid,action,row,summary):s.add(AIMissionEvent(mission_id=mid,stage="Artifact Agent",action=action,status=row.status,evidence_count=row.evidence_count,result_summary=f"Artifact {row.id} v{row.version}: {summary}"))
